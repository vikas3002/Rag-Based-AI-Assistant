import os, re, glob, hashlib
import chromadb
from pypdf import PdfReader
from docx import Document

SUPPORTED = (".pdf", ".docx", ".txt", ".md")
CHUNK_VERSION = "v2"   # change this whenever chunking logic changes -> files re-index automatically


def _docx_sections(path):
    """Split a DOCX into (page, header, body) using Title/Heading styles.
    header looks like 'Tiger — Coat and identification' so every chunk
    carries the animal/topic name, even when the body text never repeats it."""
    d = Document(path)
    out, buf = [], []
    title, section = "", ""

    def flush():
        body = "\n".join(buf).strip()
        if body:
            header = " — ".join(x for x in (title, section) if x)
            out.append((0, header, body))   # page 0 = page info not available
        buf.clear()

    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = p.style.name or ""
        if style == "Title" or style == "Heading 1":
            flush(); title, section = txt, ""
        elif style.startswith("Heading"):
            flush(); section = txt
        else:
            buf.append(txt)
    flush()

    table_rows = []
    for t in d.tables:
        for row in t.rows:
            table_rows.append(" | ".join(c.text.strip() for c in row.cells))
    if table_rows:
        out.append((0, "Tables", "\n".join(table_rows)))
    return out


def read_sections(path):
    """Returns list of (page, header, body)."""
    p = path.lower()
    if p.endswith(".pdf"):
        return [(i + 1, "", pg.extract_text() or "")
                for i, pg in enumerate(PdfReader(path).pages)]
    if p.endswith(".docx"):
        return _docx_sections(path)
    with open(path, encoding="utf-8", errors="ignore") as f:
        return [(0, "", f.read())]


def chunk_text(text, size=900, overlap=150):
    """Fixed-size windows, but short texts stay as one chunk and cuts
    prefer a sentence/line boundary instead of mid-word."""
    text = text.strip()
    if len(text) <= size:
        return [text] if text else []
    out, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            cut = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
            if cut > start + size // 2:
                end = cut + 1
        out.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in out if c]


def _hash(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return f"{h.hexdigest()}-{CHUNK_VERSION}"


class RagStore:
    def __init__(self, persist_dir, embed_docs, embed_query=None,
                 model_name="bge-small-en-v1.5"):
        self.embed_docs = embed_docs
        self.embed_query = embed_query or embed_docs
        client = chromadb.PersistentClient(path=persist_dir)
        # model name is part of the collection name, so models never get mixed
        name = re.sub(r"[^A-Za-z0-9._-]", "-", f"docs_{model_name}")
        self.col = client.get_or_create_collection(
            name=name, metadata={"hnsw:space": "cosine"})

    def count(self):
        return self.col.count()

    def remove_source(self, name):
        self.col.delete(where={"source": name})

    def index_file(self, path):
        """Returns number of chunks added (0 if unchanged or empty)."""
        name = os.path.basename(path)
        fhash = _hash(path)
        old = self.col.get(where={"source": name}, limit=1, include=["metadatas"])
        if old["ids"] and old["metadatas"][0].get("file_hash") == fhash:
            return 0                              # unchanged, skip
        self.remove_source(name)                  # drop stale chunks

        doc_id = hashlib.sha1(name.encode()).hexdigest()[:10]
        texts, metas = [], []
        for page, header, body in read_sections(path):
            for c in chunk_text(body):
                texts.append(f"{header}\n{c}" if header else c)
                metas.append({"source": name, "page": page, "doc_id": doc_id,
                              "chunk_index": len(texts) - 1, "file_hash": fhash,
                              "section": header})
        if not texts:
            return 0
        for i in range(0, len(texts), 256):       # batches
            t, m = texts[i:i + 256], metas[i:i + 256]
            self.col.add(ids=[f"{doc_id}-{x['chunk_index']}" for x in m],
                         documents=t, metadatas=m,
                         embeddings=self.embed_docs(t))
        return len(texts)

    def sync_folder(self, folder):
        os.makedirs(folder, exist_ok=True)
        files = {os.path.basename(p): p for p in glob.glob(os.path.join(folder, "*"))
                 if p.lower().endswith(SUPPORTED) and not os.path.basename(p).startswith("~$")}
        indexed = {m["source"] for m in self.col.get(include=["metadatas"])["metadatas"]}
        for gone in indexed - set(files):         # file deleted from folder
            self.remove_source(gone)
        for name, path in files.items():
            try:
                n = self.index_file(path)
                print(f"{name}: {n} new chunks" if n else f"{name}: unchanged/no text")
            except Exception as e:
                print(f"Skipped {name}: {e}")

    def retrieve(self, query, k=4, min_score=0.35):
        total = self.col.count()
        if total == 0 or not query.strip():
            return []
        res = self.col.query(query_embeddings=self.embed_query([query]),
                             n_results=min(k, total),
                             include=["documents", "metadatas", "distances"])
        hits = []
        for doc, meta, dist in zip(res["documents"][0], res["metadatas"][0],
                                   res["distances"][0]):
            score = 1 - dist                      # cosine distance -> similarity
            if score >= min_score:
                hits.append({"text": doc, "source": meta["source"],
                             "page": meta.get("page"), "score": score})
        return hits